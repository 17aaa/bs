#!/usr/bin/env python3
"""生成论文表格的 .docx 文件，可直接在 WPS/Word 中打开使用"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

def set_cell_font(cell, text, font_name='宋体', font_size=Pt(10.5), bold=False):
    """设置单元格字体"""
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = font_size
    run.bold = bold

def add_table_with_title(doc, title, headers, rows):
    """添加带标题的表格"""
    # 添加表格标题
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.font.size = Pt(10.5)
    run.bold = True

    # 创建表格
    num_cols = len(headers)
    num_rows = len(rows) + 1  # +1 for header
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    # 设置表头
    for j, header in enumerate(headers):
        set_cell_font(table.rows[0].cells[j], header, bold=True)

    # 设置数据行
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            set_cell_font(table.rows[i + 1].cells[j], str(cell_text))

    # 添加空行分隔
    doc.add_paragraph()

def main():
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(10.5)

    # 标题页
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('毕业论文表格')
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.font.size = Pt(18)
    run.bold = True

    doc.add_paragraph('以下为论文中所有表格，可直接复制到论文文档中使用。')
    doc.add_paragraph()

    # ========== 表3-1 用户表 ==========
    add_table_with_title(doc, '表3-1 用户表（users）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['username', 'varchar(50)', '是', '用户名'],
            ['password', 'varchar(255)', '是', '密码（加密存储）'],
            ['email', 'varchar(100)', '是', '邮箱'],
            ['phone', 'varchar(20)', '是', '手机号'],
            ['avatar', 'varchar(255)', '是', '头像 URL'],
            ['school', 'varchar(100)', '是', '学校'],
            ['major', 'varchar(100)', '是', '专业'],
            ['bio', 'text', '是', '个人简介'],
            ['status', 'tinyint(1)', '否', '状态：1 正常，2 禁用'],
            ['ban_reason', 'varchar(500)', '是', '封禁原因'],
            ['banned_at', 'timestamp', '是', '封禁时间'],
            ['platform_balance', 'decimal(65,0)', '否', '平台积分余额（Wei精度），用于NFT交易支付和粉丝代币公募，注册时默认赠送100积分'],
            ['created_at', 'datetime', '否', '创建时间'],
            ['updated_at', 'datetime', '否', '更新时间'],
        ])

    # ========== 表3-2 用户角色表 ==========
    add_table_with_title(doc, '表3-2 用户角色表（user_roles）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['user_id', 'bigint(20)', '否', '用户 ID'],
            ['role', 'varchar(20)', '否', '角色：user / admin / super_admin'],
            ['created_at', 'datetime', '否', '创建时间'],
            ['updated_at', 'datetime', '否', '更新时间'],
        ])

    # ========== 表3-3 钱包表 ==========
    add_table_with_title(doc, '表3-3 钱包表（wallets）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['user_id', 'bigint(20)', '否', '用户 ID，唯一'],
            ['wallet_address', 'varchar(42)', '否', '钱包地址，唯一'],
            ['keystore_path', 'varchar(200)', '是', '密钥文件路径'],
            ['is_bound', 'tinyint(1)', '否', '是否已绑定，默认0'],
            ['bound_at', 'timestamp', '是', '绑定时间'],
            ['created_at', 'datetime', '否', '创建时间'],
            ['updated_at', 'datetime', '否', '更新时间'],
        ])

    # ========== 表3-4 NFT资产表 ==========
    add_table_with_title(doc, '表3-4 NFT资产表（nft_assets）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['token_id', 'bigint(20)', '否', 'NFT Token ID'],
            ['contract_address', 'varchar(42)', '否', '合约地址'],
            ['name', 'varchar(200)', '否', '作品名称'],
            ['description', 'text', '是', '作品描述'],
            ['category', 'varchar(50)', '是', '作品分类'],
            ['image_url', 'varchar(500)', '是', '作品图片 URL'],
            ['current_metadata_hash', 'varchar(100)', '否', '当前元数据CID'],
            ['current_version', 'int', '否', '当前版本号，默认1'],
            ['owner_address', 'varchar(42)', '否', '所有者钱包地址'],
            ['creator_address', 'varchar(42)', '否', '创作者钱包地址'],
            ['royalty_fee', 'int', '否', '版税比例（万分比），默认500'],
            ['royalty_recipient', 'varchar(42)', '是', '版税接收地址'],
            ['review_status', 'varchar(20)', '否', '审核状态：pending / approved / rejected'],
            ['status', 'tinyint(1)', '否', '状态：1 正常，2 下架，3 冻结'],
            ['created_at', 'datetime', '否', '创建时间'],
            ['updated_at', 'datetime', '否', '更新时间'],
        ])

    # ========== 表3-5 版本历史表 ==========
    add_table_with_title(doc, '表3-5 版本历史表（nft_versions）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['nft_asset_id', 'bigint(20)', '否', 'NFT资产 ID'],
            ['version', 'int', '否', '版本号'],
            ['metadata_hash', 'varchar(100)', '否', '元数据CID'],
            ['ipfs_uri', 'varchar(200)', '是', 'IPFS 完整 URI'],
            ['change_description', 'text', '是', '变更说明'],
            ['updater_address', 'varchar(42)', '否', '更新者地址'],
            ['tx_hash', 'varchar(66)', '是', '交易哈希'],
            ['block_number', 'bigint(20)', '是', '区块号'],
            ['created_at', 'datetime', '否', '创建时间'],
        ])

    # ========== 表3-6 NFT审核表 ==========
    add_table_with_title(doc, '表3-6 NFT审核表（nft_reviews）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['nft_asset_id', 'bigint(20)', '否', 'NFT资产 ID'],
            ['reviewer_id', 'bigint(20)', '否', '审核人 ID'],
            ['status', 'varchar(20)', '否', '审核状态：pending / approved / rejected'],
            ['comment', 'text', '是', '审核意见'],
            ['reject_reason', 'varchar(50)', '是', '驳回原因：inappropriate / copyright / spam / other'],
            ['reviewed_at', 'datetime', '是', '审核时间'],
            ['created_at', 'datetime', '否', '提交时间'],
        ])

    # ========== 表3-7 铸造任务表 ==========
    add_table_with_title(doc, '表3-7 铸造任务表（mint_tasks）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['task_id', 'varchar(50)', '否', '任务编号'],
            ['user_id', 'bigint(20)', '否', '用户 ID'],
            ['creator_address', 'varchar(42)', '否', '创作者地址'],
            ['name', 'varchar(200)', '否', '作品名称'],
            ['description', 'text', '是', '作品描述'],
            ['category', 'varchar(50)', '是', '作品分类'],
            ['image_url', 'varchar(500)', '是', '作品图片 URL'],
            ['metadata_url', 'varchar(500)', '是', '元数据 URL'],
            ['royalty_fee', 'int', '是', '版税比例（万分比）'],
            ['current_step', 'int', '否', '当前步骤：1 上传 2 元数据 3 铸造 4 完成'],
            ['step_status', 'varchar(20)', '否', '步骤状态：pending / processing / completed / failed'],
            ['progress_percent', 'int', '否', '进度百分比（0-100）'],
            ['error_message', 'text', '是', '错误信息'],
            ['nft_asset_id', 'bigint(20)', '是', '关联NFT资产 ID'],
            ['tx_hash', 'varchar(66)', '是', '交易哈希'],
            ['rollback_status', 'varchar(20)', '是', '回滚状态：none / pending / completed / failed'],
            ['retry_count', 'int', '是', '重试次数，最多3次'],
            ['created_at', 'datetime', '否', '创建时间'],
            ['updated_at', 'datetime', '否', '更新时间'],
        ])

    # ========== 表3-8 交易订单表 ==========
    add_table_with_title(doc, '表3-8 交易订单表（market_orders）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['order_id', 'varchar(64)', '否', '订单编号，唯一'],
            ['sale_id', 'bigint(20)', '是', '链上销售 ID'],
            ['seller_address', 'varchar(42)', '否', '卖家地址'],
            ['buyer_address', 'varchar(42)', '是', '买家地址'],
            ['nft_contract', 'varchar(42)', '否', 'NFT合约地址'],
            ['token_id', 'bigint(20)', '否', 'Token ID'],
            ['nft_asset_id', 'bigint(20)', '是', 'NFT资产 ID'],
            ['order_type', 'tinyint(1)', '否', '订单类型：1 固定价格 2 荷兰拍卖 3 出价'],
            ['price', 'bigint(20)', '否', '价格（Wei）'],
            ['start_price', 'bigint(20)', '是', '拍卖起始价格'],
            ['reserve_price', 'bigint(20)', '是', '拍卖保留价格'],
            ['payment_token', 'varchar(42)', '是', '支付代币地址'],
            ['end_time', 'bigint(20)', '是', '到期时间戳'],
            ['royalty_fee', 'int', '是', '版税比例'],
            ['royalty_recipient', 'varchar(42)', '是', '版税接收地址'],
            ['final_price', 'bigint(20)', '是', '最终成交价格'],
            ['status', 'tinyint(1)', '否', '状态：1 活跃 2 已售 3 已取消 4 已过期'],
            ['tx_hash', 'varchar(66)', '是', '交易哈希'],
            ['created_at', 'datetime', '否', '创建时间'],
            ['updated_at', 'datetime', '否', '更新时间'],
            ['completed_at', 'datetime', '是', '完成时间'],
        ])

    # ========== 表3-9 粉丝代币表 ==========
    add_table_with_title(doc, '表3-9 粉丝代币表（fan_tokens）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['token_address', 'varchar(42)', '否', '代币合约地址，唯一'],
            ['project_id', 'bigint(20)', '否', '关联项目 ID'],
            ['creator_address', 'varchar(42)', '否', '创作者地址'],
            ['name', 'varchar(100)', '否', '代币名称'],
            ['symbol', 'varchar(20)', '否', '代币符号'],
            ['total_supply', 'bigint(20)', '否', '总供应量'],
            ['public_sale_remaining', 'bigint(20)', '是', '公开销售剩余量'],
            ['public_sale_price', 'bigint(20)', '是', '公开销售价格'],
            ['public_sale_active', 'tinyint(1)', '否', '是否开启公开销售，默认0'],
            ['contract_address', 'varchar(42)', '是', '合约地址'],
            ['status', 'tinyint(1)', '否', '状态：1 正常 2 暂停 3 结束'],
            ['created_at', 'datetime', '否', '创建时间'],
        ])

    # ========== 表3-10 质押记录表 ==========
    add_table_with_title(doc, '表3-10 质押记录表（stake_records）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['user_address', 'varchar(42)', '否', '用户地址'],
            ['token_address', 'varchar(42)', '否', '代币地址'],
            ['staked_amount', 'bigint(20)', '否', '质押数量'],
            ['reward_earned', 'bigint(20)', '否', '累计奖励，默认0'],
            ['balance', 'bigint(20)', '否', '余额，默认0'],
            ['last_update_time', 'bigint(20)', '是', '最后更新时间'],
            ['created_at', 'datetime', '否', '创建时间'],
            ['updated_at', 'datetime', '否', '更新时间'],
        ])

    # ========== 表3-11 IPFS文件记录表 ==========
    add_table_with_title(doc, '表3-11 IPFS文件记录表（ipfs_file_record）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['cid', 'varchar(100)', '否', 'IPFS CID，唯一'],
            ['filename', 'varchar(255)', '是', '文件名'],
            ['file_type', 'varchar(50)', '是', '文件类型'],
            ['file_size', 'bigint(20)', '否', '文件大小，默认0'],
            ['mime_type', 'varchar(100)', '是', 'MIME 类型'],
            ['pinned', 'tinyint(1)', '否', '是否已固定，默认0'],
            ['pinned_at', 'timestamp', '是', '固定时间'],
            ['user_id', 'bigint(20)', '是', '上传用户 ID'],
            ['owner_address', 'varchar(42)', '是', '所有者地址'],
            ['nft_asset_id', 'bigint(20)', '是', '关联NFT资产 ID'],
            ['description', 'text', '是', '文件描述'],
            ['tags', 'varchar(500)', '是', '标签'],
            ['access_count', 'int', '否', '访问次数，默认0'],
            ['last_accessed_at', 'timestamp', '是', '最后访问时间'],
            ['status', 'varchar(20)', '否', '状态：active / archived'],
            ['deleted', 'tinyint(1)', '否', '逻辑删除，默认0'],
            ['created_at', 'datetime', '否', '创建时间'],
            ['updated_at', 'datetime', '否', '更新时间'],
        ])

    # ========== 表3-12 交易通知表 ==========
    add_table_with_title(doc, '表3-12 交易通知表（transaction_notifications）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['user_id', 'bigint(20)', '否', '用户 ID'],
            ['type', 'varchar(20)', '否', '通知类型：purchase / sale / bid / transfer / royalty'],
            ['order_id', 'varchar(64)', '是', '关联订单 ID'],
            ['nft_asset_id', 'bigint(20)', '是', 'NFT资产 ID'],
            ['nft_name', 'varchar(200)', '是', 'NFT名称'],
            ['nft_image_url', 'varchar(500)', '是', 'NFT图片 URL'],
            ['amount', 'bigint(20)', '是', '金额（Wei）'],
            ['counterparty_address', 'varchar(42)', '是', '交易对方地址'],
            ['tx_hash', 'varchar(66)', '是', '交易哈希'],
            ['is_read', 'tinyint(1)', '否', '是否已读，默认0'],
            ['created_at', 'datetime', '否', '创建时间'],
        ])

    # ========== 表3-13 操作审计日志表 ==========
    add_table_with_title(doc, '表3-13 操作审计日志表（audit_logs）',
        ['字段名', '数据类型', '是否为空', '说明'],
        [
            ['id', 'bigint(20)', '否', '主键、自增'],
            ['user_id', 'bigint(20)', '否', '操作用户 ID'],
            ['username', 'varchar(50)', '是', '操作用户名'],
            ['action', 'varchar(30)', '否', '操作类型：create / update / delete / login / logout / transfer / mint / buy / sell / approve / reject / ban / unban'],
            ['module', 'varchar(20)', '否', '所属模块：user / nft / market / auth / admin / wallet'],
            ['description', 'text', '是', '操作描述'],
            ['resource_type', 'varchar(30)', '是', '资源类型'],
            ['resource_id', 'varchar(50)', '是', '资源 ID'],
            ['ip_address', 'varchar(50)', '是', 'IP 地址'],
            ['request_method', 'varchar(10)', '是', '请求方法'],
            ['request_path', 'varchar(200)', '是', '请求路径'],
            ['status', 'varchar(10)', '否', '执行状态：success / failed'],
            ['error_message', 'text', '是', '错误信息'],
            ['created_at', 'datetime', '否', '创建时间'],
        ])

    # ========== 表4-1 系统开发环境 ==========
    add_table_with_title(doc, '表4-1 系统开发环境',
        ['工具或环境', '版本'],
        [
            ['操作系统', 'macOS Sonoma 14 / Windows 11'],
            ['Java', '17'],
            ['IntelliJ IDEA', '2024.1'],
            ['VS Code', '1.90'],
            ['React', '19'],
            ['Vite', '8'],
            ['MySQL', '8.0'],
            ['Redis', '7'],
            ['Polygon网络', 'Amoy 测试网'],
            ['Hardhat', '2.19'],
            ['Solidity', '0.8.20'],
            ['Git', '2.39'],
        ])

    # ========== 表5-1 主要功能测试结果 ==========
    add_table_with_title(doc, '表5-1 主要功能测试结果',
        ['测试模块', '测试内容', '预期结果', '实际结果'],
        [
            ['钱包认证', '连接MetaMask钱包', '成功获取钱包地址并显示账户信息', '通过'],
            ['钱包认证', '钱包签名登录', '成功完成签名校验并返回JWT令牌', '通过'],
            ['钱包管理', '创建服务端钱包', '成功生成钱包地址、助记词和密钥文件', '通过'],
            ['钱包管理', '助记词恢复钱包', '成功通过助记词恢复钱包地址', '通过'],
            ['文件上传', '上传作品文件', '文件成功上传并返回CID或存储记录', '通过'],
            ['文件上传', '批量上传文件', '多个文件成功上传并返回批量CID记录', '通过'],
            ['NFT铸造', '提交作品信息并铸造', '成功生成NFT记录并返回交易结果', '通过'],
            ['NFT铸造', '查看铸造进度', '实时显示四步进度和百分比', '通过'],
            ['NFT铸造', '铸造失败重试', '失败后可重试，最多3次', '通过'],
            ['NFT审核', '提交作品审核', '作品进入待审核队列', '通过'],
            ['NFT审核', '管理员审核通过', '作品审核状态更新为已通过', '通过'],
            ['NFT审核', '管理员审核驳回', '作品审核状态更新为已驳回，记录驳回原因', '通过'],
            ['市场交易', '固定价格上架', '成功创建订单并展示在市场列表中', '通过'],
            ['市场交易', '荷兰式拍卖上架', '成功创建拍卖订单，价格随时间递减', '通过'],
            ['市场交易', '购买NFT', '成功完成交易并更新订单状态', '通过'],
            ['市场交易', '取消订单', '成功取消活跃订单并归还NFT', '通过'],
            ['版本管理', '新增作品版本', '成功写入版本记录并展示版本历史', '通过'],
            ['粉丝代币', '发行代币', '成功创建代币记录（平台积分模式）', '通过'],
            ['粉丝代币', '公开销售购买', '成功购买代币并更新剩余量', '通过'],
            ['粉丝代币', '质押与奖励', '成功质押代币并计算累计奖励', '通过'],
            ['通知', '接收交易通知', '交易完成后收到对应类型通知', '通过'],
            ['通知', '标记已读', '成功标记单条或全部通知为已读', '通过'],
            ['审计', '操作日志记录', '关键操作被记录到审计日志表', '通过'],
            ['后台管理', '数据概览', '成功展示用户、NFT、交易等统计指标', '通过'],
            ['后台管理', '用户管理（封禁/解封/角色变更）', '成功执行用户管理操作', '通过'],
            ['后台管理', 'NFT管理（下架/冻结/审核）', '成功执行NFT管理操作', '通过'],
            ['后台管理', '订单统计', '成功展示订单统计数据', '通过'],
            ['后台管理', '代币销售控制', '成功激活/暂停/终止代币销售', '通过'],
        ])

    # ========== 表5-2 异常与边界测试结果 ==========
    add_table_with_title(doc, '表5-2 异常与边界测试结果',
        ['测试模块', '测试内容', '测试输入', '预期结果', '实际结果'],
        [
            ['钱包认证', '过期challenge签名', '使用超过5分钟的nonce签名', '验证失败，返回"challenge不存在或已过期"', '通过'],
            ['钱包认证', '重复签名攻击', '同一challenge两次签名验证', '第二次验证失败（challenge已删除）', '通过'],
            ['NFT铸造', 'IPFS上传失败降级', '关闭IPFS服务后铸造NFT', '降级为MinIO存储，铸造仍可完成', '通过'],
            ['NFT铸造', '铸造重试超过上限', '模拟连续3次失败后重试', '第4次重试被拒绝，提示"已达到最大重试次数"', '通过'],
            ['市场交易', '余额不足购买', '平台积分余额低于NFT价格', '交易失败，返回"平台积分余额不足"', '通过'],
            ['市场交易', '荷兰拍卖到期', '拍卖已过endTime后购买', '价格降至reservePrice，或订单自动过期', '通过'],
            ['市场交易', '卖家购买自己的NFT', '卖家地址与买家地址相同', '交易失败，返回"卖家不能购买自己的NFT"', '通过'],
            ['粉丝代币', '质押超过余额', '质押数量大于可用余额', '质押失败，返回"余额不足"', '通过'],
            ['粉丝代币', '公募购买超过剩余量', '购买数量超过publicSaleRemaining', '购买失败，返回"公募代币不足"', '通过'],
            ['文件上传', '超大文件上传', '上传超过50MB的文件', '上传失败，返回"文件大小超过限制"', '通过'],
            ['文件上传', '不支持的文件类型', '上传.exe文件', '上传失败，返回"不支持的文件类型"', '通过'],
        ])

    # ========== 表5-3 性能测试结果 ==========
    add_table_with_title(doc, '表5-3 性能测试结果',
        ['测试接口', '并发数', '平均响应时间(ms)', '95%响应时间(ms)', 'QPS', '错误率'],
        [
            ['GET /api/nft/list', '10', '45', '78', '218', '0%'],
            ['GET /api/nft/list', '50', '89', '156', '542', '0%'],
            ['GET /api/nft/list', '100', '178', '312', '556', '0%'],
            ['POST /api/wallet/challenge', '10', '12', '25', '812', '0%'],
            ['POST /api/wallet/challenge', '50', '28', '52', '1768', '0%'],
            ['POST /api/market/orders', '10', '35', '68', '282', '0%'],
            ['POST /api/market/orders', '50', '82', '145', '604', '0%'],
            ['POST /api/nft/mint-task/create', '5', '156', '245', '32', '0%'],
        ])

    # 保存文件
    output_path = '/Users/zhm/Desktop/user-service/毕业论文表格.docx'
    doc.save(output_path)
    print(f'已生成: {output_path}')

if __name__ == '__main__':
    main()